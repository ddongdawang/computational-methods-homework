# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"C:\Users\User\Desktop\学习\计算方法\计算方法_作业01_答案.docx"

doc = Document()

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')


def set_run(run, size=12, bold=False, italic=False, ascii_font='Times New Roman', east_font='宋体'):
    run.font.name = ascii_font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), east_font)


def para(text, size=12, bold=False, italic=False, align=None, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, italic=italic)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def heading(text, level=1):
    if level == 0:
        return para(text, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)
    if level == 1:
        return para(text, size=14, bold=True, space_after=8)
    return para(text, size=12.5, bold=True, space_after=6)


def set_cell_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '000000')
        borders.append(el)
    tblPr.append(borders)


def mat(data, note=None, vline_col=None):
    """矩阵用 [] 文本形式表示，等宽字体右对齐"""
    if note:
        para(note, italic=True, space_after=4)
    ncols = len(data[0])
    widths = [0] * ncols
    for row in data:
        for j, v in enumerate(row):
            widths[j] = max(widths[j], len(v))
    for row in data:
        parts = []
        for j, v in enumerate(row):
            if vline_col is not None and j == vline_col:
                parts.append('|')
            parts.append(v.rjust(widths[j]))
        line = '[' + '  '.join(parts) + ']'
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        r = p.add_run(line)
        set_run(r, size=11, ascii_font='Consolas', east_font='宋体')
    para('', size=6, space_after=2)


def summary_table(data):
    t = doc.add_table(rows=len(data), cols=len(data[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_borders(t)
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.width = Inches(1.2)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            r = p.add_run(val)
            set_run(r, size=11)
    para('', size=6, space_after=2)


def code_block(lines):
    for ln in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.left_indent = Inches(0.2)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), 'F5F5F5')
        pPr.append(shd)
        r = p.add_run(ln)
        set_run(r, size=10.5, ascii_font='Consolas', east_font='宋体')
    para('', size=6, space_after=4)


# ================= 文档正文 =================
heading('计算方法 作业01 答案', 0)

# ---------- 第1题 ----------
heading('第1题  有效数字位数')

para('判定准则：设准确值 x 的近似值为 x*，将 x* 写成 x* = 0.a1a2…×10^m（a1≠0）的形式，'
     '若绝对误差 |x* − x| ≤ 0.5×10^(m−n)，则 x* 具有 n 位有效数字。'
     '即：误差不超过第 n 位有效数字所在数位的半个单位。')

para('(1) x1* = 3.7869，准确值 x = 3.78695：', bold=True)
para('|x1* − x| = |3.7869 − 3.78695| = 0.00005 = 0.5×10⁻⁴。'
     '3.7869 = 0.37869×10¹，m = 1，由 0.5×10^(1−n) = 0.5×10⁻⁴ 得 n = 5。'
     '故 x1* 有 5 位有效数字。')

para('(2) x2* = 3.7870，准确值 x = 3.78695：', bold=True)
para('|x2* − x| = 0.00005 = 0.5×10⁻⁴，同 (1)，n = 5。故 x2* 有 5 位有效数字。')

para('(3) y1* = 9.9999，准确值 y = 10：', bold=True)
para('|y1* − y| = 0.0001 = 10⁻⁴。9.9999 = 0.99999×10¹，m = 1。'
     '0.5×10^(1−4) = 0.5×10⁻³ = 0.0005 ≥ 0.0001，而 0.5×10^(1−5) = 0.00005 < 0.0001。'
     '故 y1* 有 4 位有效数字。')

para('(4) y2* = 10.1，准确值 y = 10：', bold=True)
para('|y2* − y| = 0.1 = 10⁻¹。10.1 = 0.101×10²，m = 2。'
     '0.5×10^(2−2) = 0.5 ≥ 0.1，而 0.5×10^(2−3) = 0.05 < 0.1。'
     '故 y2* 有 2 位有效数字。')

para('(5) y3* = 10.0001，准确值 y = 10：', bold=True)
para('|y3* − y| = 0.0001 = 10⁻⁴。10.0001 = 0.100001×10²，m = 2。'
     '0.5×10^(2−5) = 0.5×10⁻³ = 0.0005 ≥ 0.0001，而 0.5×10^(2−6) = 0.00005 < 0.0001。'
     '故 y3* 有 5 位有效数字。')

para('综上：', bold=True)
summary_table([
    ['近似值', 'x1*', 'x2*', 'y1*', 'y2*', 'y3*'],
    ['有效数字位数', '5 位', '5 位', '4 位', '2 位', '5 位'],
])

# ---------- 第2题 ----------
heading('第2题  Gauss 消去法')

para('方程组 Ax = b，其中：')
mat([
    ['1', '2', '1', '-2'],
    ['2', '5', '3', '-2'],
    ['-2', '-2', '3', '5'],
    ['1', '3', '2', '5'],
], note='系数矩阵 A')

para('右端向量 b = (-1, 3, 15, 9)^T，增广矩阵 [A | b]：')
mat([
    ['1', '2', '1', '-2', '-1'],
    ['2', '5', '3', '-2', '3'],
    ['-2', '-2', '3', '5', '15'],
    ['1', '3', '2', '5', '9'],
], vline_col=4)

para('（1）矩阵左乘实现 Gauss 消去法（顺序消元）', bold=True)

para('第 1 步：主元 a11 = 1，消去第 1 列下方元素（行2 − 2×行1，行3 + 2×行1，行4 − 行1）：')
mat([
    ['1', '2', '1', '-2', '-1'],
    ['0', '1', '1', '2', '5'],
    ['0', '2', '5', '1', '13'],
    ['0', '1', '1', '7', '10'],
], vline_col=4)

para('第 2 步：主元 a22 = 1，消去第 2 列下方元素（行3 − 2×行2，行4 − 行2）：')
mat([
    ['1', '2', '1', '-2', '-1'],
    ['0', '1', '1', '2', '5'],
    ['0', '0', '3', '-3', '3'],
    ['0', '0', '0', '5', '5'],
], vline_col=4)

para('第 3 步：主元 a33 = 3，其下方已为 0，无需消元。此时系数矩阵已化为上三角矩阵。')

para('回代：', bold=True)
para('5x4 = 5  →  x4 = 1')
para('3x3 − 3x4 = 3  →  x3 = 2')
para('x2 + x3 + 2x4 = 5  →  x2 = 1')
para('x1 + 2x2 + x3 − 2x4 = −1  →  x1 = −3')
para('解：x = (x1, x2, x3, x4) = (-3, 1, 2, 1)。', bold=True)

para('（2）Gauss 列主元素消去法', bold=True)

para('第 1 步：第 1 列元素为 1, 2, -2, 1，绝对值最大者为 2（第 2 行），交换第 1、2 行：')
mat([
    ['2', '5', '3', '-2', '3'],
    ['1', '2', '1', '-2', '-1'],
    ['-2', '-2', '3', '5', '15'],
    ['1', '3', '2', '5', '9'],
], vline_col=4)

para('以 a11 = 2 为主元消元（行2 − (1/2)行1，行3 + 行1，行4 − (1/2)行1）：')
mat([
    ['2', '5', '3', '-2', '3'],
    ['0', '-1/2', '-1/2', '-1', '-5/2'],
    ['0', '3', '6', '3', '18'],
    ['0', '1/2', '1/2', '6', '15/2'],
], vline_col=4)

para('第 2 步：第 2 列（第 2 行起）元素为 -1/2, 3, 1/2，绝对值最大者为 3（第 3 行），交换第 2、3 行：')
mat([
    ['2', '5', '3', '-2', '3'],
    ['0', '3', '6', '3', '18'],
    ['0', '-1/2', '-1/2', '-1', '-5/2'],
    ['0', '1/2', '1/2', '6', '15/2'],
], vline_col=4)

para('以 a22 = 3 为主元消元（行3 + (1/6)行2，行4 − (1/6)行2）：')
mat([
    ['2', '5', '3', '-2', '3'],
    ['0', '3', '6', '3', '18'],
    ['0', '0', '1/2', '-1/2', '1/2'],
    ['0', '0', '-1/2', '11/2', '9/2'],
], vline_col=4)

para('第 3 步：第 3 列（第 3 行起）元素为 1/2, -1/2，绝对值最大者取 1/2（第 3 行），不交换。'
     '以 a33 = 1/2 为主元消元（行4 + 行3）：')
mat([
    ['2', '5', '3', '-2', '3'],
    ['0', '3', '6', '3', '18'],
    ['0', '0', '1/2', '-1/2', '1/2'],
    ['0', '0', '0', '5', '5'],
], vline_col=4)

para('回代：', bold=True)
para('5x4 = 5  →  x4 = 1')
para('(1/2)x3 − (1/2)x4 = 1/2  →  x3 = 2')
para('3x2 + 6x3 + 3x4 = 18  →  x2 = 1')
para('2x1 + 5x2 + 3x3 − 2x4 = 3  →  x1 = −3')
para('解：x = (x1, x2, x3, x4) = (-3, 1, 2, 1)，与（1）结果一致。', bold=True)

# ---------- 第3题 ----------
heading('第3题  Cholesky 分解')

para('A 为 3×3 Hilbert 矩阵，对称正定：')
mat([
    ['1', '1/2', '1/3'],
    ['1/2', '1/3', '1/4'],
    ['1/3', '1/4', '1/5'],
], note='矩阵 A')

para('Cholesky 分解：A = L·L^T，其中 L 为下三角矩阵（对角元为正）。')
para('递推公式（i ≥ j）：', bold=True)
para('l_ii = √( a_ii − Σ_{k=1}^{i−1} l_ik² )')
para('l_ij = ( a_ij − Σ_{k=1}^{j−1} l_ik·l_jk ) / l_jj   (i > j)')

para('计算过程：', bold=True)
para('l11 = √a11 = 1')
para('l21 = a21 / l11 = 1/2，  l31 = a31 / l11 = 1/3')
para('l22 = √(a22 − l21²) = √(1/3 − 1/4) = √(1/12) = 1/√12')
para('l32 = (a32 − l31·l21) / l22 = (1/4 − 1/6) / (1/√12) = 1/√12')
para('l33 = √(a33 − l31² − l32²) = √(1/5 − 1/9 − 1/12) = √(1/180) = 1/√180')

para('结果：', bold=True)
mat([
    ['1', '0', '0'],
    ['1/2', '1/√12', '0'],
    ['1/3', '1/√12', '1/√180'],
], note='下三角矩阵 L')

para('（等价地，有理化后：l22 = l32 = √3/6 ≈ 0.2887，l33 = √5/30 ≈ 0.0745。）')
para('验证：计算 L·L^T 可还原为 A，故分解正确。', italic=True)

# ---------- 第4题 ----------
heading('第4题  追赶法（Crout 法）')

para('三对角方程组 Ax = f：')
mat([
    ['2', '1', '0', '0'],
    ['1', '4', '1', '0'],
    ['0', '1', '4', '1'],
    ['0', '0', '1', '2'],
], note='三对角矩阵 A')

para('f = (-3, 6, 14, -2)^T')

para('记主对角线 b = (2, 4, 4, 2)，下对角线 a = (1, 1, 1)，上对角线 c = (1, 1, 1)。')

para('追赶法（Crout 分解）将 A 分解为 A = LU：', bold=True)
para('L 为下三角矩阵，主对角元素记为 p_i，副对角元素为 a_i（即原下对角线）；')
para('U 为单位上三角矩阵，副对角元素记为 q_i。')

para('递推公式：p1 = b1，q_i = c_i / p_i，p_i = b_i − a_i·q_(i−1)（i = 2, 3, 4）。')

para('计算 L、U 的系数：', bold=True)
para('p1 = 2，  q1 = c1/p1 = 1/2')
para('p2 = b2 − a2·q1 = 4 − 1/2 = 7/2，  q2 = c2/p2 = 2/7')
para('p3 = b3 − a3·q2 = 4 − 2/7 = 26/7，  q3 = c3/p3 = 7/26')
para('p4 = b4 − a4·q3 = 2 − 7/26 = 45/26')
para('即 p = (2, 7/2, 26/7, 45/26)，q = (1/2, 2/7, 7/26)。')

para('“追”——前推解 Ly = f：', bold=True)
para('2y1 = -3  →  y1 = -3/2')
para('y1 + (7/2)y2 = 6  →  y2 = 15/7')
para('y2 + (26/7)y3 = 14  →  y3 = 83/26')
para('y3 + (45/26)y4 = -2  →  y4 = -3')

para('“赶”——回代解 Ux = y：', bold=True)
para('x4 = y4 = -3')
para('x3 + (7/26)x4 = 83/26  →  x3 = 4')
para('x2 + (2/7)x3 = 15/7  →  x2 = 1')
para('x1 + (1/2)x2 = -3/2  →  x1 = -2')

para('解：x = (x1, x2, x3, x4) = (-2, 1, 4, -3)。', bold=True)
para('验证：代入原方程组逐行成立，结果正确。', italic=True)

# ---------- 第5题 ----------
heading('第5题  MATLAB 程序 myDoolittle')

para('完整程序：', bold=True)
code_block([
    'function [L, U] = myDoolittle(A)',
    '% myDoolittle  用 Doolittle 法实现矩阵 A 的 LU 分解',
    '%   输入：A —— n×n 可逆矩阵',
    '%   输出：L —— 单位下三角矩阵（对角元全为 1）',
    '%         U —— 上三角矩阵',
    '%   满足：A = L*U',
    'n = size(A, 1);',
    'L = eye(n);     % 单位下三角矩阵',
    'U = zeros(n);   % 上三角矩阵',
    '',
    'for k = 1:n',
    '    % (1) 先求 U 的第 k 行',
    '    for j = k:n',
    '        U(k, j) = A(k, j) - L(k, 1:k-1) * U(1:k-1, j);',
    '    end',
    '    % (2) 再求 L 的第 k 列',
    '    for i = k+1:n',
    '        L(i, k) = (A(i, k) - L(i, 1:k-1) * U(1:k-1, k)) / U(k, k);',
    '    end',
    'end',
    'end',
])

para('测试：', bold=True)
code_block([
    'A = [1 -1 3; 2 -4 6; 4 -9 2];',
    '[L, U] = myDoolittle(A);',
    'disp(L), disp(U), disp(L*U)',
])

para('运行输出：', bold=True)
code_block([
    'L =',
    '     1     0     0',
    '     2     1     0',
    '     4   2.5     1',
    'U =',
    '     1    -1     3',
    '     0    -2     0',
    '     0     0   -10',
    'L*U =',
    '     1    -1     3',
    '     2    -4     6',
    '     4    -9     2',
])

para('可见 L*U = A，程序输出正确。', bold=True)

doc.save(OUT)
print('saved:', OUT)
